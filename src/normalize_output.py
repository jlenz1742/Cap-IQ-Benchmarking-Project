#!/usr/bin/env python3
"""
Read a Capital IQ workbook (produced by generate_template.py, then
refreshed and saved in Excel with the CapIQ Add-in) and reshape the
'Income Statement' sheet into a tidy long-format CSV:

    CompanyName, Identifier, Industry, LineItem, Year, Value

'Year' here is the year-offset label from the template (e.g. "FY (latest)",
"FY-1", "FY-2", ...) -- each company's own most recent reported fiscal
year and the years before it, not necessarily the same calendar year
across companies.

Usage:
    python src/normalize_output.py \
        --workbook output/CapIQ_Income_Statement_Template.xlsx \
        --out output/income_statement_tidy.csv
"""
import argparse
from pathlib import Path

from openpyxl import load_workbook


def normalize(workbook_path, sheet_name="Income Statement"):
    wb = load_workbook(workbook_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise SystemExit(f"Sheet '{sheet_name}' not found in {workbook_path}")
    ws = wb[sheet_name]

    header = [c.value for c in ws[1]]
    year_cols = list(range(6, len(header) + 1))  # columns F.. are year columns
    year_labels = [header[c - 1] for c in year_cols]

    rows = []
    for r in range(2, ws.max_row + 1):
        company = ws.cell(row=r, column=1).value
        if company is None:
            continue
        identifier = ws.cell(row=r, column=2).value
        industry = ws.cell(row=r, column=3).value
        line_item = ws.cell(row=r, column=4).value
        for col, year_label in zip(year_cols, year_labels):
            value = ws.cell(row=r, column=col).value
            rows.append({
                "CompanyName": company,
                "Identifier": identifier,
                "Industry": industry,
                "LineItem": line_item,
                "Year": year_label,
                "Value": value,
            })
    return rows


def write_csv(rows, out_path):
    import csv
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["CompanyName", "Identifier", "Industry", "LineItem", "Year", "Value"])
        writer.writeheader()
        writer.writerows(rows)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workbook", default="output/CapIQ_Income_Statement_Template.xlsx")
    ap.add_argument("--sheet", default="Income Statement")
    ap.add_argument("--out", default="output/income_statement_tidy.csv")
    args = ap.parse_args()

    rows = normalize(args.workbook, args.sheet)
    write_csv(rows, args.out)

    n_errors = sum(1 for r in rows if isinstance(r["Value"], str) and r["Value"].startswith("#"))
    print(f"Wrote {len(rows)} rows to {args.out}")
    if n_errors:
        print(f"Warning: {n_errors} cells contain formula errors (e.g. #N/A) -- "
              f"check identifiers/mnemonics for those rows.")


if __name__ == "__main__":
    main()
